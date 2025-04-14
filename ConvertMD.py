import flet as ft
from flet import Colors, Icons, FontWeight, ScrollMode
import markdown2
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, RGBColor
from pathlib import Path
import traceback
import yaml
from datetime import datetime
import pandas as pd
import os

def parse_metadata(md_content):
    metadata = {}
    cleaned_content = md_content
    if md_content.startswith('---'):
        parts = md_content.split('---', 2)
        if len(parts) >= 3:
            yaml_content = parts[1].strip()
            try:
                metadata = yaml.safe_load(yaml_content) or {}
            except yaml.YAMLError:
                pass
            cleaned_content = parts[2].lstrip('\n')
    return metadata, cleaned_content

def format_metadata(metadata):
    if not metadata:
        return ""
    formatted = []
    for key, value in metadata.items():
        if isinstance(value, list):
            value = ", ".join(str(v) for v in value)
        elif isinstance(value, datetime):
            value = value.strftime('%d.%m.%Y, %H:%M')
        formatted.append(f"{key}: {value}")
    return "\n".join(formatted)

def markdown_to_html(markdown_text):
    try:
        return markdown2.markdown(markdown_text, extras=["tables", "fenced-code-blocks"])
    except Exception as e:
        raise Exception(f"Ошибка преобразования Markdown: {str(e)}")

def html_to_word(html, output_file, metadata):
    try:
        doc = Document()
        styles = doc.styles
        font_name = 'Arial'

        if 'Metadata' not in styles:
            metadata_style = styles.add_style('Metadata', WD_STYLE_TYPE.PARAGRAPH)
            metadata_style.font.name = font_name
            metadata_style.font.size = Pt(10)
            metadata_style.font.color.rgb = RGBColor(100, 100, 100)
            metadata_style.paragraph_format.space_after = Pt(12)

        if 'Heading' not in styles:
            heading_style = styles.add_style('Heading', WD_STYLE_TYPE.PARAGRAPH)
            heading_style.font.name = font_name
            heading_style.font.size = Pt(14)
            heading_style.font.bold = True
            heading_style.font.color.rgb = RGBColor(0, 0, 0)

        if metadata:
            metadata_text = format_metadata(metadata)
            p = doc.add_paragraph(metadata_text, style='Metadata')
            p.paragraph_format.space_after = Pt(12)

        soup = BeautifulSoup(html, 'html.parser')
        current_list_type = None

        def process_list(element, level=0):
            nonlocal current_list_type
            list_type = "ul" if element.name == "ul" else "ol"
            prev_list_type = current_list_type
            current_list_type = list_type
            for child in element.children:
                if child.name == "li":
                    text = child.get_text().strip()
                    if text:
                        p = doc.add_paragraph(style='ListBullet' if list_type == "ul" else 'ListNumber')
                        p.text = text
                        nested_lists = child.find_all(['ul', 'ol'])
                        for nested in nested_lists:
                            process_list(nested, level + 1)
                elif child.name in ['ul', 'ol']:
                    process_list(child, level + 1)
            current_list_type = prev_list_type

        for element in soup.find_all(True):
            if element.name in ["h1", "h2", "h3"]:
                doc.add_paragraph(element.get_text(), style='Heading')
            elif element.name == "p":
                doc.add_paragraph(element.get_text())
            elif element.name == "table":
                try:
                    rows = element.find_all("tr")
                    if not rows:
                        continue
                    cols = len(rows[0].find_all(["th", "td"]))
                    table = doc.add_table(rows=1, cols=cols)
                    hdr_cells = table.rows[0].cells
                    for i, header in enumerate(rows[0].find_all(["th", "td"])):
                        hdr_cells[i].text = header.get_text().strip()
                    for row in rows[1:]:
                        row_cells = row.find_all("td")
                        if len(row_cells) != cols:
                            continue
                        new_row = table.add_row().cells
                        for i, cell in enumerate(row_cells):
                            new_row[i].text = cell.get_text().strip()
                except Exception as e:
                    print(f"Ошибка обработки таблицы: {str(e)}")
                    continue
            elif element.name in ["ul", "ol"]:
                process_list(element)

        doc.save(output_file)
        return True, f"Файл сохранён: {output_file}"
    except Exception as e:
        return False, f"Ошибка создания Word: {str(e)}"

def extract_h2_sections(md_content):
    lines = md_content.splitlines()
    data = {}
    current_h2 = None
    buffer = []
    for line in lines:
        if line.startswith("## "):
            if current_h2:
                data[current_h2] = "\n".join(buffer).strip()
            current_h2 = line[3:].strip()
            buffer = []
        elif current_h2:
            buffer.append(line)
    if current_h2:
        data[current_h2] = "\n".join(buffer).strip()
    return data

def markdowns_to_excel(md_files, output_file):
    all_headers = set()
    file_data = []
    for md_file in md_files:
        with open(md_file, 'r', encoding='utf-8') as f:
            md_content = f.read()
        _, content = parse_metadata(md_content)
        sections = extract_h2_sections(content)
        all_headers.update(sections.keys())
        file_data.append(sections)
    all_headers = sorted(all_headers)
    rows = []
    for sections in file_data:
        row = [sections.get(header, "") for header in all_headers]
        rows.append(row)
    df = pd.DataFrame(rows, columns=all_headers)
    df.insert(0, "Filename", [f.name for f in md_files])
    df.to_excel(output_file, index=False)

def main(page: ft.Page):
    page.title = "MD to Word/Excel Converter"
    page.window_width = 800
    page.window_height = 600

    selected_files = []

    def show_snackbar(message, color=Colors.RED_400):
        page.snack_bar = ft.SnackBar(content=ft.Text(message), bgcolor=color)
        page.snack_bar.open = True
        page.update()

    def on_files_selected(e):
        if e.files:
            selected_files.clear()
            selected_files.extend(Path(f.path) for f in e.files)
            source_file_input.value = ", ".join(f.name for f in selected_files)
            page.update()

    def select_output_dir(e):
        output_dir_picker.get_directory_path()

    def on_dir_selected(e):
        if e.path:
            output_dir_input.value = e.path
            page.update()

    def convert_files(e):
        if not selected_files:
            show_snackbar("Не выбраны .md файлы!")
            return
        output_dir = output_dir_input.value or str(selected_files[0].parent)
        Path(output_dir).mkdir(parents=True, exist_ok=True)
        result_text.value = ""
        page.update()
        if format_selector.value == "word":
            success_count = 0
            for md_file in selected_files:
                try:
                    with open(md_file, 'r', encoding='utf-8') as f:
                        md_content = f.read()
                    metadata, cleaned_content = parse_metadata(md_content)
                    html = markdown_to_html(cleaned_content)
                    output_file = Path(output_dir) / f"{md_file.stem}.docx"
                    success, message = html_to_word(html, str(output_file), metadata)
                    if success:
                        result_text.value += f"✓ Успешно: {md_file.name} → {output_file.name}\n"
                        success_count += 1
                    else:
                        result_text.value += f"✗ Ошибка: {md_file.name} - {message}\n"
                except Exception as e:
                    result_text.value += f"✗ Ошибка обработки {md_file.name}: {str(e)}\n"
                page.update()
            show_snackbar(
                f"Готово! Обработано {success_count}/{len(selected_files)} файлов",
                Colors.GREEN_400 if success_count == len(selected_files) else Colors.ORANGE_400
            )
        else:
            output_file = Path(output_dir) / "md_summary.xlsx"
            try:
                markdowns_to_excel(selected_files, output_file)
                result_text.value = f"✓ Успешно: {output_file.name} создан\n"
            except Exception as err:
                result_text.value = f"✗ Ошибка Excel: {str(err)}\n"
            page.update()

    source_file_picker = ft.FilePicker(on_result=on_files_selected)
    output_dir_picker = ft.FilePicker(on_result=on_dir_selected)
    page.overlay.extend([source_file_picker, output_dir_picker])

    user_documents = str(Path.home() / "Documents")

    source_file_input = ft.TextField(label="Выбранные .md файлы", expand=True, read_only=True)
    output_dir_input = ft.TextField(label="Целевая директория (по умолчанию - Документы)", expand=True, value=user_documents)
    result_text = ft.TextField(multiline=True, min_lines=10, expand=True, read_only=True)
    format_selector = ft.RadioGroup(
        content=ft.Row([
            ft.Radio(value="word", label="DOCX"),
            ft.Radio(value="excel", label="Excel"),
        ]),
        value="word"
    )

    page.add(
        ft.Column([
            ft.Text("Конвертер Markdown → DOCX/Excel", size=20, weight=FontWeight.BOLD),
            ft.Row([
                source_file_input,
                ft.IconButton(icon=Icons.ATTACH_FILE, on_click=lambda _: source_file_picker.pick_files(allow_multiple=True, allowed_extensions=["md"]), tooltip="Выбрать .md файлы")
            ]),
            ft.Row([
                output_dir_input,
                ft.IconButton(icon=Icons.FOLDER_OPEN_OUTLINED, on_click=select_output_dir, tooltip="Выбрать целевую директорию")
            ]),
            ft.Text("Формат вывода:"),
            format_selector,
            ft.ElevatedButton("Конвертировать", on_click=convert_files, style=ft.ButtonStyle(padding=20)),
            ft.Text("Результат:"),
            result_text
        ], scroll=ScrollMode.ALWAYS, expand=True)
    )

if __name__ == "__main__":
    ft.app(target=main)